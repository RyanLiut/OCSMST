import datetime as dt
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

TARGET = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
RESULT_ROOT = Path(__file__).resolve().parent
FIRST_COL = 3
BLOCK_WIDTH = 10
BLOCKS = [('Top 1', 1, 'value'), ('Top 5 (mean)', 5, 'mean'), ('Top 5 (max)', 5, 'max'),
          ('Top 10 (mean)', 10, 'mean'), ('Top 10 (max)', 10, 'max')]
METRICS = ['coverage', 'productivity', 'deg_mean', 'deg_std', 'weight_sum',
           'acc', 'prec', 'recall', 'F1', 'total_time_s']
NORM_LABEL = '纯共现_norm'
FIRST_LABEL = '纯共现（NAACL）'
J_NORM, D_NORM = 'Jaccard only', 'Dice only'
J_NONORM, D_NONORM = 'Jaccard only（无norm）', 'Dice only（无norm）'
VARIANTS = {'norm': RESULT_ROOT / 'pure_similarity',
            'nonorm': RESULT_ROOT / 'pure_similarity_nonorm'}
LABEL_SPEC = [(J_NORM, 'Jaccard_only', 'norm'), (D_NORM, 'Dice_only', 'norm'),
              (J_NONORM, 'Jaccard_only_nonorm', 'nonorm'),
              (D_NONORM, 'Dice_only_nonorm', 'nonorm')]
# 每组 (锚点行方法名, 要新增的行)：锚点行必须已存在，新行插在锚点行之后。
ANCHOR_GROUPS = [(NORM_LABEL, (J_NORM, D_NORM)), ('Dice only', (J_NONORM, D_NONORM))]
ORDER = [FIRST_LABEL, NORM_LABEL] + [label for label, _, _ in LABEL_SPEC] + \
        ['共现_blend_Jaccard_0.1', '共现_blend_Dice_0.2', '共现_blend_Emb_0.4',
         'Prompt_only', 'Skill+Codex']
EXCLUDED = {'case6', 'case12'}
CASES = [f'case{i}' for i in range(1, 13)]
DECIMALS = 10
N_COLS = FIRST_COL + len(BLOCKS) * BLOCK_WIDTH - 1


def preflight():
    if TARGET.with_name(f'~${TARGET.name}').exists():
        raise SystemExit(f'存在 WPS 锁文件，先关闭：{TARGET.name}')
    out = subprocess.run(['lsof', str(TARGET)], capture_output=True, text=True)
    if out.stdout.strip():
        raise SystemExit(f'目标文件被占用，先关闭：\n{out.stdout}')


def row_values(ws, row):
    return [ws.cell(row, col).value for col in range(1, N_COLS + 1)]


def as_number(value):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return None
    rounded = round(float(value), DECIMALS)
    return int(rounded) if rounded.is_integer() and abs(rounded) < 1e15 else rounded


def scan_blocks(ws):
    blocks = []
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 2).value == FIRST_LABEL:
            blocks.append({'start': row, 'label': ws.cell(row, 1).value})
    for block in blocks:
        labels, row = {}, block['start']
        while row <= ws.max_row and ws.cell(row, 2).value is not None and \
                (row == block['start'] or ws.cell(row, 1).value is None):
            labels.setdefault(ws.cell(row, 2).value, row)
            row += 1
        block['labels'] = labels
        block['end'] = row - 1
    return blocks


def load_summaries():
    summaries = {}
    for name, path in VARIANTS.items():
        frame = pd.read_csv(path / 'evaluation_summary.csv')
        bad = frame[frame['status'] != 'ok']
        if len(bad):
            raise SystemExit(f'{name} summary 存在 {len(bad)} 行非 ok 状态')
        summaries[name] = frame
    return summaries


def build_expected(summaries):
    expected = {}
    for label, method, variant in LABEL_SPEC:
        summary = summaries[variant]
        for case in CASES:
            for bi, (_, top_k, statistic) in enumerate(BLOCKS):
                rows = summary[(summary['case'] == case) & (summary['method'] == method)
                               & (summary['top_k'] == top_k) & (summary['statistic'] == statistic)]
                if len(rows) != 1:
                    raise SystemExit(f'{case}/{method}/B{bi} 匹配到 {len(rows)} 行')
                row = rows.iloc[0]
                for mi, metric in enumerate(METRICS):
                    expected[(case, label, bi, mi)] = as_number(row[metric])
        kept = summary[~summary['case'].isin(EXCLUDED)]
        for bi, (_, top_k, statistic) in enumerate(BLOCKS):
            rows = kept[(kept['method'] == method) & (kept['top_k'] == top_k)
                        & (kept['statistic'] == statistic)]
            if len(rows) != 10:
                raise SystemExit(f'平均块 {method}/B{bi} 匹配到 {len(rows)} 行')
            for mi, metric in enumerate(METRICS):
                expected[('avg', label, bi, mi)] = as_number(rows[metric].mean())
    return expected


def insert_rows(ws, blocks):
    pending = []
    for block in blocks:
        for anchor, labels in ANCHOR_GROUPS:
            missing = tuple(label for label in labels if label not in block['labels'])
            if not missing:
                continue
            if anchor not in block['labels']:
                raise SystemExit(f"case {block['label']} 缺少锚点行 {anchor}")
            pending.append((block['labels'][anchor], missing, block['label']))
    row_map = {row: row for row in range(1, ws.max_row + 1)}
    added = []
    for anchor_row, missing, case_label in sorted(pending, key=lambda item: item[0], reverse=True):
        style = [ws.cell(anchor_row, col)._style for col in range(1, N_COLS + 1)]
        ws.insert_rows(anchor_row + 1, len(missing))
        for offset, label in enumerate(missing, 1):
            new_row = anchor_row + offset
            for col in range(1, N_COLS + 1):
                ws.cell(new_row, col)._style = style[col - 1]
            ws.cell(new_row, 2).value = label
        row_map = {old: (pos if pos < anchor_row + 1 else pos + len(missing))
                   for old, pos in row_map.items()}
        added.append((case_label, anchor_row + 1, list(missing)))
    return row_map, added, sum(len(missing) for _, missing, _ in pending)


def fill(ws, blocks, expected):
    written = {}
    for block in blocks:
        case = 'avg' if isinstance(block['label'], str) else f"case{block['label']}"
        for label, _, _ in LABEL_SPEC:
            row = block['labels'][label]
            if ws.cell(row, 2).value != label:
                raise SystemExit(f"case {block['label']} 行{row} 方法名不是 {label}")
            written[row] = {2: label}
            for bi, _ in enumerate(BLOCKS):
                for mi, _ in enumerate(METRICS):
                    col = FIRST_COL + bi * BLOCK_WIDTH + mi
                    value = expected[(case, label, bi, mi)]
                    ws.cell(row, col).value = value
                    written[row][col] = value
    return written


def merge_case_columns(ws, blocks):
    for block in blocks:
        ws.merge_cells(start_row=block['start'], start_column=1,
                       end_row=block['end'], end_column=1)


def verify(path, backup, before, row_map, blocks, written, expected, row_delta):
    problems = []
    old_ws = load_workbook(backup)['主实验']
    new_ws = load_workbook(path)['主实验']

    for old_row, values in before.items():
        new_row = row_map[old_row]
        now = row_values(new_ws, new_row)
        for col, (a, b) in enumerate(zip(values, now), 1):
            want = written.get(new_row, {}).get(col, a)
            if want != b:
                problems.append(f'行{old_row}->新行{new_row} 列{col}: {b!r} != {want!r}')

    inserted = set(range(1, new_ws.max_row + 1)) - set(row_map.values())
    if len(inserted) != row_delta:
        problems.append(f'新增行数 {len(inserted)} != {row_delta}')
    for row in sorted(inserted):
        for col in range(1, N_COLS + 1):
            want = written.get(row, {}).get(col)
            got = new_ws.cell(row, col).value
            if want != got:
                problems.append(f'新行{row} 列{col}: {got!r} != {want!r}')

    old_merges = {str(r) for r in old_ws.merged_cells.ranges}
    old_block = set()
    for merge in old_merges:
        match = re.fullmatch(r'A(\d+):A(\d+)', merge)
        if match and int(match.group(1)) >= 3:
            old_block.add(merge)
    expected_merges = (old_merges - old_block) | \
        {f'A{b["start"]}:A{b["end"]}' for b in blocks}
    got_merges = {str(r) for r in new_ws.merged_cells.ranges}
    if expected_merges != got_merges:
        problems.append(f'合并区不符: 缺 {sorted(expected_merges - got_merges)} '
                        f'多 {sorted(got_merges - expected_merges)}')

    if new_ws.max_row != old_ws.max_row + row_delta:
        problems.append(f'行数 {new_ws.max_row} != {old_ws.max_row + row_delta}')

    for block in blocks:
        labels = [new_ws.cell(row, 2).value for row in range(block['start'], block['end'] + 1)]
        if labels != ORDER:
            problems.append(f"case {block['label']} 方法顺序: {labels}")
        if len(labels) != len(ORDER):
            problems.append(f"case {block['label']} 行数 {len(labels)} != {len(ORDER)}")
        case = 'avg' if isinstance(block['label'], str) else f"case{block['label']}"
        template = block['labels'][NORM_LABEL]
        for label, _, _ in LABEL_SPEC:
            row = block['labels'][label]
            for bi, _ in enumerate(BLOCKS):
                for mi, _ in enumerate(METRICS):
                    col = FIRST_COL + bi * BLOCK_WIDTH + mi
                    want = expected[(case, label, bi, mi)]
                    got = new_ws.cell(row, col).value
                    if got != want:
                        problems.append(f'{case}/{label}/B{bi}M{mi}: {got!r} != {want!r}')
                    if new_ws.cell(row, col).number_format != \
                            new_ws.cell(template, col).number_format:
                        problems.append(f'{case}/{label}/列{col} 数字格式不一致')
    return problems


def main():
    preflight()
    summaries = load_summaries()
    expected = build_expected(summaries)
    tmp = TARGET.with_name(f'{TARGET.stem}.tmp{TARGET.suffix}')
    tmp.unlink(missing_ok=True)
    stamp = dt.datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = TARGET.with_name(f'{TARGET.stem}.bak-{stamp}{TARGET.suffix}')
    shutil.copy2(TARGET, backup)
    print(f'备份: {backup.name}', flush=True)

    wb = load_workbook(TARGET)
    ws = wb['主实验']
    before = {row: row_values(ws, row) for row in range(1, ws.max_row + 1)}
    for rng in list(ws.merged_cells.ranges):
        if rng.min_col == 1 and rng.max_col == 1 and rng.min_row >= 3:
            ws.unmerge_cells(str(rng))

    blocks = scan_blocks(ws)
    if len(blocks) != 13:
        raise SystemExit(f'期望 13 个 case 块，实际 {len(blocks)}')
    row_map, added, row_delta = insert_rows(ws, blocks)
    print(f'插入行: {added}', flush=True)

    blocks = scan_blocks(ws)
    written = fill(ws, blocks, expected)
    merge_case_columns(ws, blocks)

    wb.save(tmp)
    problems = verify(tmp, backup, before, row_map, blocks, written, expected, row_delta)
    if problems:
        for item in problems[:40]:
            print('FAIL', item, flush=True)
        tmp.unlink(missing_ok=True)
        raise SystemExit(f'校验失败 {len(problems)} 项，未覆盖原文件')

    os.replace(tmp, TARGET)
    print(f'已写入: {TARGET}', flush=True)
    print('\n填入的 Top-1 / Top-10(mean) F1:', flush=True)
    for case in CASES + ['avg']:
        parts = []
        for label, _, _ in LABEL_SPEC:
            parts.append(f'{label}={expected[(case, label, 0, 8)]}'
                         f'/{expected[(case, label, 3, 8)]}')
        print(f'  {case:6s} ' + '  '.join(parts))
    return 0


if __name__ == '__main__':
    sys.exit(main())
