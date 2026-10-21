"""在「主实验」每个块的 共现_blend_Emb_0.4 行后插入 共现_blend_Emb_0.4（无norm） 行。

数据来源：ARR_Oct/main_evaluation/no_norm_emb/evaluation_summary.csv
（main.py evaluate_method(normalized=False) + summarize，已与 app 10/10 case 核对）

注意事项：
- openpyxl 的 insert_rows 不调整合并区 → 先取消全部合并，插入后按块重算恢复；
- 从后往前插块，坐标始终用原表坐标；
- case6/case12 缺 Qwen3 相似度矩阵 → 新行留空（与 norm 行一致）；
- 平均行 = 10 个 case（1,2,3,4,5,7,8,9,10,11）算术均值。
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SRC = ROOT / 'ARR_Oct/main_evaluation/no_norm_emb/evaluation_summary.csv'
SHEET = '主实验'
LABEL = '共现_blend_Emb_0.4（无norm）'
ANCHOR = '共现_blend_Emb_0.4'
BLOCK_STARTS = [3, 14, 25, 36, 47, 58, 69, 80, 91, 102, 113, 124, 135]
CASE_IDS = [f'case{i}' for i in range(1, 13)] + ['平均（除去case6和12）']
INCLUDE_IDS = [f'case{i}' for i in (1, 2, 3, 4, 5, 7, 8, 9, 10, 11)]
METRIC_COLS = [('coverage', 'coverage'), ('productivity', 'productivity'),
               ('deg_mean', 'deg_mean'), ('deg_std', 'deg_std'),
               ('weight_sum', 'weight_sum'), ('acc', 'acc'), ('prec', 'prec'),
               ('recall', 'recall'), ('F1', 'F1'), ('time', 'total_time_s')]
BLOCK_SPECS = [(1, 'value'), (5, 'mean'), (5, 'max'), (10, 'mean'), (10, 'max')]
HEADER_MERGES = ['A1:A2', 'B1:B2', 'C1:L1', 'M1:V1', 'W1:AF1', 'AG1:AP1', 'AQ1:AZ1']


def cell_tuple(ws, row):
    return tuple(ws.cell(row, c).value for c in range(3, 53))


def snapshot(ws, offsets, nrows):
    """{(case_id, label): (colA, 50 个值)} 用于前后比对；nrows=每块行数（11 或 12）"""
    out = {}
    for case_id, start, off in zip(CASE_IDS, BLOCK_STARTS, offsets):
        for j in range(nrows):
            row = start + off + j
            out[(case_id, ws.cell(row, 2).value)] = (ws.cell(row, 1).value, cell_tuple(ws, row))
    return out


def build_values():
    s = pd.read_csv(SRC)
    assert (s.status == 'ok').sum() == 50 and (s.status == 'missing').sum() == 10, s.status.value_counts()
    assert set(s.loc[s.status == 'missing', 'case']) == {'case6', 'case12'}
    s = s.set_index(['case', 'top_k', 'statistic'])
    values = {}
    for case_id in CASE_IDS[:12]:
        cells = []
        for k, st in BLOCK_SPECS:
            row = s.loc[(case_id, k, st)]
            for _, src_col in METRIC_COLS:
                v = row[src_col]
                cells.append(None if pd.isna(v) else float(v))
        values[case_id] = cells
    avg = []
    for idx in range(50):
        vals = [values[c][idx] for c in INCLUDE_IDS]
        assert all(v is not None for v in vals), f'平均值第 {idx} 格缺数据'
        avg.append(float(np.mean(vals)))
    values[CASE_IDS[12]] = avg
    return values


def preflight():
    if XLSX.with_name(f'~${XLSX.name}').exists():
        sys.exit(f'文件被占用（存在锁文件 ~${XLSX.name}）：请先关闭 WPS/Excel')
    holding = subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip()
    if holding:
        sys.exit(f'文件被占用：\n{holding}\n请先关闭 WPS/Excel')
    p = subprocess.run(['lsof', '+D', str(XLSX.parent)], capture_output=True, text=True).stdout
    if any(XLSX.name in line for line in p.splitlines()):
        sys.exit(f'文件被占用：\n{p}')
    print('预检通过：无进程占用、无锁文件')


def main():
    preflight()
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = XLSX.with_name(f'{XLSX.stem}.bak-{stamp}{XLSX.suffix}')
    shutil.copy2(XLSX, backup)
    print('备份 →', backup)

    values = build_values()
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    before = snapshot(ws, [0] * 13, 11)
    before_merges = {str(m) for m in ws.merged_cells.ranges}
    before_others = {name: [tuple(ws2.cell(r, c).value for c in range(1, ws2.max_column + 1))
                            for r in range(1, ws2.max_row + 1)]
                     for name, ws2 in ((n, wb[n]) for n in wb.sheetnames if n != SHEET)}
    assert ws.max_row == 145 and ws.max_column == 52, (ws.max_row, ws.max_column)

    for rng in sorted(before_merges, key=len, reverse=True):
        ws.unmerge_cells(rng)

    for i in range(len(BLOCK_STARTS) - 1, -1, -1):
        target = BLOCK_STARTS[i] + 9
        ws.insert_rows(target)
        for c in range(1, 53):
            ws.cell(target, c)._style = copy(ws.cell(target - 1, c)._style)
        ws.cell(target, 2).value = LABEL
        written = 0
        for idx, v in enumerate(values[CASE_IDS[i]]):
            if v is None:
                continue
            v = round(v, 10)
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            ws.cell(target, 3 + idx).value = v
            written += 1
        print(f'{CASE_IDS[i]:>16s} ← 第 {target} 行，写入 {written}/50 格')

    ws.merge_cells('A1:A2')
    for rng in HEADER_MERGES[1:]:
        ws.merge_cells(rng)
    for i, start in enumerate(BLOCK_STARTS):
        ws.merge_cells(f'A{start + i}:A{start + i + 11}')
    wb.save(XLSX)
    print('已保存 →', XLSX)

    # ---------- 校验 ----------
    new_wb = load_workbook(XLSX)
    nws = new_wb[SHEET]
    errs = []
    if (nws.max_row, nws.max_column) != (158, 52):
        errs.append(f'尺寸 {(nws.max_row, nws.max_column)} != (158, 52)')
    expect_merges = set(HEADER_MERGES) | {f'A{BLOCK_STARTS[i] + i}:A{BLOCK_STARTS[i] + i + 11}'
                                          for i in range(len(BLOCK_STARTS))}
    got_merges = {str(m) for m in nws.merged_cells.ranges}
    if got_merges != expect_merges:
        errs.append(f'合并区不符：多 {sorted(got_merges - expect_merges)}、缺 {sorted(expect_merges - got_merges)}')
    for name, rows_before in before_others.items():
        ws2 = new_wb[name]
        rows_now = [tuple(ws2.cell(r, c).value for c in range(1, ws2.max_column + 1))
                    for r in range(1, ws2.max_row + 1)]
        if rows_before != rows_now:
            errs.append(f'其他 sheet「{name}」发生变化')

    after = snapshot(nws, list(range(13)), 12)
    for key, payload in before.items():
        if key not in after:
            errs.append(f'原行丢失：{key}')
        elif after[key] != payload:
            diff = [(j, a, b) for j, (a, b) in enumerate(zip(payload[1], after[key][1])) if a != b]
            errs.append(f'原行被改动：{key}，差异 {diff[:5]}')
    extra = {k for k in after if k not in before}
    if {k[1] for k in extra} != {LABEL} or len(extra) != 13:
        errs.append(f'新增行不符：{sorted(k[1] for k in extra)}')

    for i, case_id in enumerate(CASE_IDS):
        row = BLOCK_STARTS[i] + i + 9
        if nws.cell(row, 2).value != LABEL:
            errs.append(f'{case_id} 新行标签位置错：{nws.cell(row, 2).value}')
            continue
        for idx, v in enumerate(values[case_id]):
            if v is None:
                continue
            want = round(v, 10)
            if isinstance(want, float) and want.is_integer():
                want = int(want)
            got = nws.cell(row, 3 + idx).value
            if got is None or abs(float(got) - float(want)) > 0:
                errs.append(f'{case_id} 新行第 {idx} 格 {got} != {want}')
        # 均值独立复算：由 10 个 case 的新行反推
        if case_id == '平均（除去case6和12）':
            case_rows = {c: BLOCK_STARTS[j] + j + 9 for j, c in enumerate(CASE_IDS)}
            for idx in range(50):
                src = [nws.cell(case_rows[c], 3 + idx).value for c in INCLUDE_IDS]
                if any(v is None for v in src):
                    errs.append(f'平均格 {idx} 有 None')
                elif abs(float(nws.cell(row, 3 + idx).value) - float(np.mean(src))) > 1e-9:
                    errs.append(f'平均格 {idx} 复算不符')
        if i < 12:  # 新行样式应与上一行（Emb norm 行）一致
            for c in (2, 3, 52):
                a, b = nws.cell(row, c)._style, nws.cell(row - 1, c)._style
                if a != b:
                    errs.append(f'{case_id} 第 {c} 列样式与 Emb 行不一致')

    print('\n校验：' + ('全部通过 ✓' if not errs else f'{len(errs)} 项失败'))
    for e in errs:
        print('  ✗', e)
    if errs:
        print('还原备份…')
        shutil.copy2(backup, XLSX)
        return 1
    print(f'新增 13 行（12 case + 平均）；case6/case12 留空；其他 4 个 sheet 零变化')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
