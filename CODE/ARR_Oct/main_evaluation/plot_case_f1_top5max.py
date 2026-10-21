"""主实验对比图：OCMST（纯共现 NAACL）vs OCSMST_m（共现_blend_Emb_0.1（无norm+minmax+merge））。

- 指标：Top 5 (max) 块的 F1（列 W:AF 中第 9 个指标，即 AE 列）；
- 每个 case 两根重叠条：后面宽条为 OCSMST_m，前面窄条为 OCMST；
- case6/case12 无 OCSMST_m 数据 → 不画；
- 排版与 λ 权重图一致（单栏 3.5 in、Times/STIX、矢量 PDF）。

输出：ARR_Oct/main_evaluation/minmax01_merge/case_f1_top5max_ocmst_vs_ocsmstm.pdf（及 png 预览）
"""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'ARR_Oct/main_evaluation/minmax01_merge'
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SHEET = '主实验'
LABEL_OURS = '共现_blend_Emb_0.1（无norm+minmax+merge）'
LABEL_BASE = '纯共现（NAACL）'
TOP5MAX_START = 23          # W 列：Top 5 (max) 块起点
F1_OFFSET = 8               # 块内第 9 个指标 = F1 → 列 31 (AE)

plt.rcParams.update({
    'font.family': 'serif',
    'font.serif': ['Times New Roman', 'STIXGeneral', 'DejaVu Serif'],
    'mathtext.fontset': 'stix',
    'font.size': 8,
    'axes.labelsize': 9,
    'xtick.labelsize': 7.5,
    'ytick.labelsize': 7.5,
    'legend.fontsize': 7.5,
    'axes.linewidth': 0.6,
    'xtick.major.width': 0.6,
    'ytick.major.width': 0.6,
    'xtick.major.size': 2.5,
    'ytick.major.size': 2.5,
    'pdf.fonttype': 42,
})


def main():
    ws = load_workbook(XLSX)[SHEET]
    f1_col = TOP5MAX_START + F1_OFFSET
    cases, base, ours = [], [], []
    for i in range(13):
        start = 3 + 15 * i
        case_id = ws.cell(start, 1).value
        if not isinstance(case_id, (int, float)):      # 跳过「平均」块
            continue
        rows = {ws.cell(start + j, 2).value: start + j for j in range(15)}
        if LABEL_OURS not in rows:
            continue
        b, o = ws.cell(rows[LABEL_BASE], f1_col).value, ws.cell(rows[LABEL_OURS], f1_col).value
        if o is None or b is None:
            continue                      # case6/case12：跳过
        cases.append(int(case_id))
        base.append(float(b))
        ours.append(float(o))
    cases, base, ours = np.array(cases), np.array(base), np.array(ours)
    orders = [1,3,4,5,6,0,7,8,9,2,]
    cases, base, ours = cases[orders], base[orders], ours[orders]
    delta = ours - base
    print('case, OCMST, OCSMST_m, Δ:', list(zip(cases, base.round(4), ours.round(4), delta.round(4))))
    print('平均: OCMST=%.4f, OCSMST_m=%.4f, Δ=%+.4f' % (base.mean(), ours.mean(), delta.mean()))

    x = np.arange(len(cases))
    fig, ax = plt.subplots(figsize=(3.5, 2.8))
    back = ax.bar(x, ours * 100, width=0.72, color='#4E79A7', edgecolor='#3A5F86',
                  linewidth=0.4, zorder=2, label='OCSMST$\\mathregular{_m}$')
    front = ax.bar(x, base * 100, width=0.42, color='#F2A65A', edgecolor='#C9822F',
                   linewidth=0.4, zorder=3, label='OCMST')

    for idx, d in enumerate(delta):
        ax.annotate(f'{d * 100:+.1f}',
                    xy=(idx, max(base[idx], ours[idx]) * 100),
                    xytext=(0, 2), textcoords='offset points',
                    ha='center', va='bottom', fontsize=6,
                    color='#4E79A7' if d > 1e-9 else ('0.40' if d < -1e-9 else '0.50'))

    ax.set_xlabel('Dataset')
    ax.set_ylabel('F1@5$^*$(%)')
    ax.set_xticks(x)
    ax.set_xticklabels(1 + np.arange(10))#cases)
    ax.set_ylim(0, 134)
    ax.set_yticks([0, 20, 40, 60, 80, 100, 120])
    ax.grid(axis='y', color='0.88', linewidth=0.5, zorder=0)
    ax.set_axisbelow(True)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.legend(handles=[front, back], loc='upper center', ncols=2, frameon=False,
              handlelength=1.2, columnspacing=1.4, borderaxespad=0.1,
              handletextpad=0.5)

    fig.tight_layout(pad=0.3)
    for ext in ('pdf', 'png'):
        fig.savefig(OUT / f'case_f1_top5max_ocmst_vs_ocsmstm.{ext}',
                    dpi=300 if ext == 'png' else None, bbox_inches='tight', facecolor='white')
    plt.close(fig)
    print('已保存 →', OUT / 'case_f1_top5max_ocmst_vs_ocsmstm.pdf')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
