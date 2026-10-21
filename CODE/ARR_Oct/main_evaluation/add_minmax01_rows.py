"""在「主实验」的 共现_blend_Emb_0.4（无norm+merge）行后插入两行新方法：

1. 共现_blend_Emb_0.1（无norm+minmax）
2. 共现_blend_Emb_0.1（无norm+minmax+merge）

数据来源：
  ARR_Oct/main_evaluation/minmax01/evaluation_summary.csv
  ARR_Oct/main_evaluation/minmax01_merge/evaluation_summary.csv
（选择协议与主实验一致：top_n=5/10、max_tied_trees=6000；cooc 先做 min-max 到 [0,1]
  再以 λ=0.1 与 Qwen3-8B 单独定义相似度融合；merge 行对每棵候选树先跑 merge_edge）

- case6/case12 缺相似度矩阵 → 留空；
- case3/case5 功能数 ≥20，productivity 由 SMM 返回 N/A（不补算）→ 该 5 格留空；
- 平均行按可用 case 求均值（productivity 为 8 个 case，其余 10 个）；
- 从后往前插块，坐标用原表坐标；先取消合并、插完按 15 行/块重算合并区。
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SRC_U = ROOT / 'ARR_Oct/main_evaluation/minmax01/evaluation_summary.csv'
SRC_M = ROOT / 'ARR_Oct/main_evaluation/minmax01_merge/evaluation_summary.csv'
SHEET = '主实验'
LABEL_U = '共现_blend_Emb_0.1（无norm+minmax）'
LABEL_M = '共现_blend_Emb_0.1（无norm+minmax+merge）'
ANCHOR = '共现_blend_Emb_0.4（无norm+merge）'
OLD_STARTS = [3 + 13 * i for i in range(13)]
CASE_IDS = [f'case{i}' for i in range(1, 13)] + ['平均（除去case6和12）']
INCLUDE_IDS = [f'case{i}' for i in (1, 2, 3, 4, 5, 7, 8, 9, 10, 11)]
METRIC_COLS = [('coverage', 'coverage'), ('productivity', 'productivity'),
               ('deg_mean', 'deg_mean'), ('deg_std', 'deg_std'),
               ('weight_sum', 'weight_sum'), ('acc', 'acc'), ('prec', 'prec'),
               ('recall', 'recall'), ('F1', 'F1'), ('time', 'total_time_s')]
BLOCK_SPECS = [(1, 'value'), (5, 'mean'), (5, 'max'), (10, 'mean'), (10, 'max')]
HEADER_MERGES = ['A1:A2', 'B1:B2', 'C1:L1', 'M1:V1', 'W1:AF1', 'AG1:AP1', 'AQ1:AZ1']


def cell_tuple(ws, row):
    return tuple(ws.cell(row, c).value for c in range(3, 53))


def snapshot(ws, starts, nrows):
    out = {}
    for case_id, start in zip(CASE_IDS, starts):
        for j in range(nrows):
            row = start + j
            out[(case_id, ws.cell(row, 2).value)] = (ws.cell(row, 1).value, cell_tuple(ws, row))
    return out


def build_values(src):
    s = pd.read_csv(src)
    assert (s.status == 'ok').sum() == 50 and (s.status == 'missing').sum() == 10, s.status.value_counts()
    assert set(s.loc[s.status == 'missing', 'case']) == {'case6', 'case12'}
    s = s.set_index(['case', 'top_k', 'statistic'])
    values = {}
    for case_id in CASE_IDS[:12]:
        cells = []
        for k, st in BLOCK_SPECS:
            row = s.loc[(case_id, k, st)]
            for _, src_col in METRIC_COLS:
                v = row[src_col]
                cells.append(None if pd.isna(v) else float(v))
        values[case_id] = cells
    avg, counts = [], []
    for idx in range(50):
        vals = [values[c][idx] for c in INCLUDE_IDS if values[c][idx] is not None]
        counts.append(len(vals))
        avg.append(float(np.mean(vals)) if vals else None)
    values[CASE_IDS[12]] = avg
    return values, counts


def preflight():
    if XLSX.with_name(f'~${XLSX.name}').exists():
        sys.exit(f'文件被占用（存在锁文件 ~${XLSX.name}）：请先关闭 WPS/Excel')
    holding = subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip()
    if holding:
        sys.exit(f'文件被占用：\n{holding}\n请先关闭 WPS/Excel')
    p = subprocess.run(['lsof', '+D', str(XLSX.parent)], capture_output=True, text=True).stdout
    if any(XLSX.name in line for line in p.splitlines()):
        sys.exit(f'文件被占用：\n{p}\n请先关闭 WPS/Excel')
    print('预检通过：无进程占用、无锁文件')


def main():
    preflight()
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = XLSX.with_name(f'{XLSX.stem}.bak-{stamp}{XLSX.suffix}')
    shutil.copy2(XLSX, backup)
    print('备份 →', backup)

    values_u, _ = build_values(SRC_U)
    values_m, _ = build_values(SRC_M)
    wb = load_workbook(XLSX)
    ws = wb[SHEET]
    before = snapshot(ws, OLD_STARTS, 13)
    before_merges = {str(m) for m in ws.merged_cells.ranges}
    expect_before = set(HEADER_MERGES) | {f'A{OLD_STARTS[i]}:A{OLD_STARTS[i] + 12}' for i in range(13)}
    assert before_merges == expect_before, sorted(before_merges ^ expect_before)
    for i, start in enumerate(OLD_STARTS):
        assert ws.cell(start + 10, 2).value == ANCHOR, (i, ws.cell(start + 10, 2).value)

    for rng in sorted(before_merges, key=len, reverse=True):
        ws.unmerge_cells(rng)

    for i in range(12, -1, -1):
        target = OLD_STARTS[i] + 11          # 无norm+merge 行的下一行
        ws.insert_rows(target, 2)
        for offset, (label, values) in enumerate(((LABEL_U, values_u), (LABEL_M, values_m))):
            row = target + offset
            for c in range(1, 53):
                ws.cell(row, c)._style = copy(ws.cell(target - 1, c)._style)
            ws.cell(row, 2).value = label
            written = 0
            for idx, v in enumerate(values[CASE_IDS[i]]):
                if v is None:
                    continue
                v = round(v, 10)
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                ws.cell(row, 3 + idx).value = v
                written += 1
            print(f'{CASE_IDS[i]:>16s} {label[-12:]:>14s} ← 第 {row} 行，{written}/50 格')

    for rng in HEADER_MERGES:
        ws.merge_cells(rng)
    for i in range(13):
        ws.merge_cells(f'A{3 + 15 * i}:A{17 + 15 * i}')
    wb.save(XLSX)
    print('已保存 →', XLSX)
    return verify(backup, values_u, values_m)


def verify(backup, values_u, values_m):
    NEW_STARTS = [3 + 15 * i for i in range(13)]
    bwb = load_workbook(backup)
    nwb = load_workbook(XLSX)
    bws, nws = bwb[SHEET], nwb[SHEET]
    errs = []

    if nws.max_row != bws.max_row + 26 or nws.max_column != 52:
        errs.append(f'尺寸 {(nws.max_row, nws.max_column)} != ({bws.max_row + 26}, 52)')
    expect_merges = set(HEADER_MERGES) | {f'A{NEW_STARTS[i]}:A{NEW_STARTS[i] + 14}' for i in range(13)}
    got_merges = {str(m) for m in nws.merged_cells.ranges}
    if got_merges != expect_merges:
        errs.append(f'合并区不符：多 {sorted(got_merges - expect_merges)}、缺 {sorted(expect_merges - got_merges)}')
    expect_backup = set(HEADER_MERGES) | {f'A{OLD_STARTS[i]}:A{OLD_STARTS[i] + 12}' for i in range(13)}
    if {str(m) for m in bws.merged_cells.ranges} != expect_backup:
        errs.append('备份合并区与 13 行/块预期不符，无法比对')

    for name in bwb.sheetnames:
        if name == SHEET:
            continue
        b2, n2 = bwb[name], nwb[name]
        rb = [tuple(b2.cell(r, c).value for c in range(1, b2.max_column + 1)) for r in range(1, b2.max_row + 1)]
        rn = [tuple(n2.cell(r, c).value for c in range(1, n2.max_column + 1)) for r in range(1, n2.max_row + 1)]
        if rb != rn:
            errs.append(f'其他 sheet「{name}」发生变化')

    before = snapshot(bws, OLD_STARTS, 13)
    after = snapshot(nws, NEW_STARTS, 15)
    for key, payload in before.items():
        if key not in after:
            errs.append(f'原行丢失：{key}')
        elif after[key] != payload:
            diff = [(j, a, b) for j, (a, b) in enumerate(zip(payload[1], after[key][1])) if a != b]
            errs.append(f'原行被改动：{key}，差异 {diff[:5]}')
    extra = {k for k in after if k not in before}
    if len(extra) != 26 or {k[1] for k in extra} != {LABEL_U, LABEL_M}:
        errs.append(f'新增行不符：{sorted(str(k[1]) for k in extra)}')

    case_rows = {}
    blanks = 0
    for i, case_id in enumerate(CASE_IDS):
        base = NEW_STARTS[i] + 11
        case_rows[case_id] = base
        for offset, (label, values) in enumerate(((LABEL_U, values_u), (LABEL_M, values_m))):
            row = base + offset
            if nws.cell(row, 2).value != label:
                errs.append(f'{case_id} 新行标签位置错：{nws.cell(row, 2).value}')
                continue
            for c in (2, 3, 52):
                if nws.cell(row, c)._style != nws.cell(row - 1, c)._style:
                    errs.append(f'{case_id} 新行第 {c} 列样式与上一行不一致')
            for idx, v in enumerate(values[case_id]):
                got = nws.cell(row, 3 + idx).value
                if v is None:
                    if got is not None:
                        errs.append(f'{case_id} 第 {idx} 格应留空，实际 {got}')
                    blanks += 1
                    continue
                want = round(v, 10)
                if isinstance(want, float) and want.is_integer():
                    want = int(want)
                if got is None or float(got) != float(want):
                    errs.append(f'{case_id} 第 {idx} 格 {got} != {want}')
    if blanks != 2 * (10 + 50 + 50):
        errs.append(f'留空格数 {blanks} != 220')

    for idx in range(50):
        src = [nws.cell(case_rows[c], 3 + idx).value for c in INCLUDE_IDS]
        src = [v for v in src if v is not None]
        want = float(np.mean(src))
        got = nws.cell(case_rows[CASE_IDS[12]], 3 + idx).value
        if got is None or abs(float(got) - want) > 1e-9:
            errs.append(f'平均格 {idx} 复算 {got} != {want}')

    print('\n校验：' + ('全部通过 ✓' if not errs else f'{len(errs)} 项失败'))
    for e in errs[:10]:
        print('  ✗', e)
    if errs:
        print('还原备份…')
        shutil.copy2(backup, XLSX)
        return 1
    print(f'新增 26 行；留空 {blanks} 格（case6/case12 各 50×2、case3/case5 productivity 各 5×2）；'
          '平均行按可用 case 计；其他 6 个 sheet 零变化')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
