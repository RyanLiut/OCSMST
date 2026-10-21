"""在「消融实验-语义融合（cooc_minmax_0.1）」表底部追加一个「平均」块。

- 结构与每个 case 的 55 行块完全同构（cooccur_only / random / char_sim /
  13 模型 × 4 综合方式的 embsim_L融合）；
- 每格 = 该组合在 10 个 case 上的 F1 均值（跳过空格，如 case5/case8 的
  单独例句（均值池化）只有 8 个 case 有值 → 按 8 个 case 平均）；
- case 列写「平均」并纵向合并；写完校验其他 sheet 与原 550 行零变化。
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SHEET = '消融实验-语义融合（cooc_minmax_0.1）'
CASE_ROWS = [(2 + 55 * i, 56 + 55 * i) for i in range(10)]   # 每个 case 块的起止行
AVG_LABEL = '平均'


def preflight():
    if XLSX.with_name(f'~${XLSX.name}').exists():
        sys.exit(f'文件被占用（存在锁文件 ~${XLSX.name}）：请先关闭 WPS/Excel')
    holding = subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip()
    if holding:
        sys.exit(f'文件被占用：\n{holding}\n请先关闭 WPS/Excel')
    print('预检通过：无进程占用、无锁文件')


def snapshot(wb, skip=()):
    out = {}
    for name in wb.sheetnames:
        if name in skip:
            continue
        ws = wb[name]
        out[name] = ([tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
                      for r in range(1, ws.max_row + 1)],
                     {str(m) for m in ws.merged_cells.ranges})
    return out


def main():
    preflight()
    if not XLSX.stat().st_mode & 0o200:
        XLSX.chmod(XLSX.stat().st_mode | 0o600)
        print('已解除只读')
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = XLSX.with_name(f'{XLSX.stem}.bak-{stamp}{XLSX.suffix}')
    shutil.copy2(XLSX, backup)
    print('备份 →', backup)

    wb = load_workbook(XLSX)
    before = snapshot(wb)
    ws = wb[SHEET]
    assert ws.max_row == 551, ws.max_row

    # 每个 case 块的标签列（B/C/D）完全一致 → 取第一块作为模板
    template = [(ws.cell(2 + j, 2).value, ws.cell(2 + j, 3).value, ws.cell(2 + j, 4).value)
                for j in range(55)]
    # 逐格求 10 个 case 的均值（跳过空格）
    averages, counts = [], []
    for j in range(55):
        vals = [ws.cell(start + j, 5).value for start, _ in CASE_ROWS]
        vals = [float(v) for v in vals if v is not None and not (isinstance(v, float) and np.isnan(v))]
        averages.append(float(np.mean(vals)) if vals else None)
        counts.append(len(vals))

    start = 552
    for j, (label_b, label_c, label_d) in enumerate(template):
        row = start + j
        ws.cell(row, 2).value = label_b
        ws.cell(row, 3).value = label_c
        ws.cell(row, 4).value = label_d
        if averages[j] is not None:
            value = round(averages[j], 10)
            ws.cell(row, 5).value = int(value) if isinstance(value, float) and value.is_integer() else value
    ws.cell(start, 1).value = AVG_LABEL
    ws.merge_cells(start_row=start, start_column=1, end_row=start + 54, end_column=1)
    ws.cell(start, 1).alignment = Alignment(vertical='center')
    wb.save(XLSX)
    print(f'已追加「平均」块：第 {start}–{start + 54} 行 → {XLSX.name}')

    # ---------- 校验 ----------
    new_wb = load_workbook(XLSX)
    nws = new_wb[SHEET]
    errs = []
    after = snapshot(new_wb, skip=(SHEET,))
    for name, payload in before.items():
        if name == SHEET:
            continue
        if after.get(name) != payload:
            errs.append(f'其他 sheet「{name}」发生变化')
    if nws.max_row != 606 or nws.max_column != 7:
        errs.append(f'尺寸 {(nws.max_row, nws.max_column)} != (606, 7)')
    expect_merges = {f'A{2 + 55 * i}:A{56 + 55 * i}' for i in range(10)} | {'A552:A606'}
    got_merges = {str(m) for m in nws.merged_cells.ranges}
    if got_merges != expect_merges:
        errs.append(f'合并区不符：多 {sorted(got_merges - expect_merges)}、缺 {sorted(expect_merges - got_merges)}')
    if nws.cell(552, 1).value != AVG_LABEL:
        errs.append(f'平均块标签错：{nws.cell(552, 1).value}')
    # 原 550 行零变化
    bws = load_workbook(backup)[SHEET]
    for r in range(1, 552):
        for c in range(1, 8):
            if bws.cell(r, c).value != nws.cell(r, c).value:
                errs.append(f'原第 {r} 行第 {c} 列被改动')
                break
    # 平均块逐格复算
    for j in range(55):
        row = 552 + j
        if (nws.cell(row, 2).value, nws.cell(row, 3).value, nws.cell(row, 4).value) != template[j]:
            errs.append(f'第 {row} 行标签与模板不符')
        vals = [nws.cell(start2 + j, 5).value for start2, _ in CASE_ROWS]
        vals = [float(v) for v in vals if v is not None]
        want = float(np.mean(vals))
        got = nws.cell(row, 5).value
        if got is None or abs(float(got) - want) > 1e-9:
            errs.append(f'第 {row} 行均值 {got} != {want}')
    print('\n校验：' + ('全部通过 ✓' if not errs else f'{len(errs)} 项失败'))
    for e in errs[:10]:
        print('  ✗', e)
    if errs:
        print('还原备份…')
        shutil.copy2(backup, XLSX)
        return 1
    pooled = sum(counts)
    print(f'平均块 55 格：{sum(1 for c in counts if c == 10)} 格按 10 case 平均，'
          f'{sum(1 for c in counts if c == 8)} 格按 8 case 平均（case5/case8 缺均值池化）')
    print('要点：cooccur_only={:.4f}，random={:.4f}，char_sim={:.4f}，embsim 52 格宏平均={:.4f}'.format(
        averages[0], averages[1], averages[2],
        float(np.nanmean([a for a in averages[3:] if a is not None]))))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
