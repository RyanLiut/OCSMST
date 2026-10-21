"""把 cooc min-max 融合权重扫描结果写入工作簿新 Sheet。

- Sheet 名：消融实验-融合权重（无norm-cooc_minmax）
- 结构与「消融实验-融合权重（无norm）」同构（case × 权重 0~1 步长 0.1），
  但只保留 Qwen3-Embedding-8B 的 Emb 融合，并与旧口径（原式加权）并列对比；
- 新口径数值来源：ARR_Oct/main_evaluation/blend_cooc_minmax/results.csv（Top-1 F1，
  Top-1 = 最大权重生成树，不做并列树 std 排序）；
- 旧口径数值来源：本工作簿「消融实验-融合权重（无norm）」表 blend_Emb (Qwen3-...) 列；
- 写入前 lsof/锁文件预检 + 时间戳备份；写完校验其他 sheet 零变化。
"""
from __future__ import annotations

import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SRC = ROOT / 'ARR_Oct/main_evaluation/blend_cooc_minmax/results.csv'
SHEET = '消融实验-融合权重（无norm-cooc_minmax）'
TEMPLATE = '消融实验-融合权重（无norm）'
CASE_IDS = [f'case{i}' for i in range(1, 13)]
WEIGHTS = [round(i / 10, 1) for i in range(11)]
NOTE = ('口径：无norm（原始共现计数、不做行归一化）；cooc 先做 off-diagonal min-max 到 '
        '[0,1]，再 blend = weight×sim + (1−weight)×cooc_minmax；Top-1 = 最大权重生成树'
        '（不做并列树 std 排序，因此权重=0 大量并列时与旧表的 Top-1 可能不同）；'
        '旧列为原式 weight×sim + (1−weight)×cooc 原始计数。')


def preflight():
    if XLSX.with_name(f'~${XLSX.name}').exists():
        sys.exit(f'文件被占用（存在锁文件 ~${XLSX.name}）：请先关闭 WPS/Excel')
    holding = subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip()
    if holding:
        sys.exit(f'文件被占用：\n{holding}\n请先关闭 WPS/Excel')
    print('预检通过：无进程占用、无锁文件')


def snapshot(wb, skip=()):
    out = {}
    for name in wb.sheetnames:
        if name in skip:
            continue
        ws = wb[name]
        out[name] = ([tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
                      for r in range(1, ws.max_row + 1)],
                     {str(m) for m in ws.merged_cells.ranges})
    return out


def main():
    preflight()
    df = pd.read_csv(SRC)
    ok = df[df.status == 'ok']
    assert len(ok) == 110 and set(ok.case) == set(CASE_IDS) - {'case6', 'case12'}, df.status.value_counts()
    wide = ok.pivot(index='case', columns='weight', values='F1')

    if not XLSX.stat().st_mode & 0o200:
        XLSX.chmod(XLSX.stat().st_mode | 0o600)
        print('已解除只读')
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = XLSX.with_name(f'{XLSX.stem}.bak-{stamp}{XLSX.suffix}')
    backup.write_bytes(XLSX.read_bytes())
    print('备份 →', backup)

    wb = load_workbook(XLSX)
    before = snapshot(wb)
    legacy_ws = wb[TEMPLATE]
    legacy = {}
    last_case = None
    for row in range(2, legacy_ws.max_row + 1):
        weight = legacy_ws.cell(row, 2).value
        if weight is None:
            continue
        if legacy_ws.cell(row, 1).value is not None:   # 合并单元格：case 只在首行
            last_case = legacy_ws.cell(row, 1).value
        value = legacy_ws.cell(row, 6).value           # F 列 = blend_Emb (Qwen3-...)
        legacy[(f'case{int(float(last_case))}', round(float(weight), 1))] = (
            None if value is None or (isinstance(value, float) and np.isnan(value)) else float(value))

    if SHEET in wb.sheetnames:
        wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET)
    headers = ['case', '权重', 'blend_Emb (Qwen3-Embdding-8B)',
               'blend_Emb (Qwen3-Embdding-8B)（cooc_minmax）']
    ws.append(headers)
    for c in range(1, 5):
        ws.cell(1, c).font = Font(name='宋体', size=11)
    for letter, width in zip('ABCD', [8, 10, 35.9, 44]):
        ws.column_dimensions[letter].width = width

    for case in CASE_IDS:
        start = ws.max_row + 1
        for w in WEIGHTS:
            new = wide.loc[case, w] if case in wide.index else np.nan
            old = legacy.get((case, w))
            ws.append([int(case.replace('case', '')), w,
                       None if old is None else round(old, 10),
                       None if pd.isna(new) else round(float(new), 10)])
        ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
        ws.cell(start, 1).alignment = Alignment(vertical='center')
    assert ws.max_row == 133, ws.max_row
    ws.cell(1, 6).value = NOTE
    ws.cell(1, 6).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['F'].width = 70
    wb.save(XLSX)
    print(f'已写入「{SHEET}」{ws.max_row} 行 × {ws.max_column} 列 → {XLSX.name}')

    # ---------- 校验 ----------
    new_wb = load_workbook(XLSX)
    errs = []
    after = snapshot(new_wb, skip=(SHEET,))
    for name, payload in before.items():
        if name == SHEET:
            continue
        if name not in after:
            errs.append(f'sheet 丢失：{name}')
        elif after[name] != payload:
            errs.append(f'sheet「{name}」内容或合并区发生变化')
    nws = new_wb[SHEET]
    if (nws.max_row, nws.max_column) != (133, 6):
        errs.append(f'新 sheet 尺寸 {(nws.max_row, nws.max_column)} != (133, 6)')
    if {str(m) for m in nws.merged_cells.ranges} != {f'A{2 + 11 * i}:A{12 + 11 * i}' for i in range(12)}:
        errs.append('新 sheet case 列合并区不符')
    for i, case in enumerate(CASE_IDS):
        for j, w in enumerate(WEIGHTS):
            row = 2 + 11 * i + j
            if j == 0 and nws.cell(row, 1).value != int(case.replace('case', '')):
                errs.append(f'第 {row} 行 case 值错位')
            if abs(nws.cell(row, 2).value - w) > 1e-12:
                errs.append(f'第 {row} 行 权重 错位')
            got_new, got_old = nws.cell(row, 4).value, nws.cell(row, 3).value
            want_new = wide.loc[case, w] if case in wide.index else np.nan
            want_old = legacy.get((case, w))
            if pd.isna(want_new):
                if got_new is not None:
                    errs.append(f'第 {row} 行 新列应为空，实际 {got_new}')
            elif got_new is None or abs(float(got_new) - float(round(want_new, 10))) > 0:
                errs.append(f'第 {row} 行 新列 {got_new} != {round(float(want_new), 10)}')
            if want_old is None:
                if got_old is not None:
                    errs.append(f'第 {row} 行 旧列应为空，实际 {got_old}')
            elif got_old is None or abs(float(got_old) - float(round(want_old, 10))) > 0:
                errs.append(f'第 {row} 行 旧列 {got_old} != {round(float(want_old), 10)}')
    print('\n校验：' + ('全部通过 ✓' if not errs else f'{len(errs)} 项失败'))
    for e in errs[:10]:
        print('  ✗', e)
    if errs:
        print('还原备份…')
        XLSX.write_bytes(backup.read_bytes())
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
