"""把 minmax λ=0.1 的「语义融合」消融结果写入工作簿新 Sheet。

- Sheet 名：消融实验-语义融合（cooc_minmax_0.1）
- 结构与「消融实验-语义融合」同构（每 case 55 行：cooccur_only / random / char_sim /
  13 模型 × 4 综合方式的 embsim_L融合），只含 case1-5、7-11（6/12 不做）；
- cooccur_only 为 λ=0 纯共现、与 minmax 无关 → 直接沿用原表数值；
- random = 5 个种子 F1_top1 均值；embsim/char_sim = Top-1 树 F1_top1；
- 数据源：ARR_Oct/main_evaluation/semantic_fusion_minmax01/results.csv；
- 写入前预检 + 时间戳备份；写完校验其他 sheet 零变化。
"""
from __future__ import annotations

import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
SRC = ROOT / 'ARR_Oct/main_evaluation/semantic_fusion_minmax01/results.csv'
SHEET = '消融实验-语义融合（cooc_minmax_0.1）'
TEMPLATE = '消融实验-语义融合'
CASES = [f'case{i}' for i in (1, 2, 3, 4, 5, 7, 8, 9, 10, 11)]
MODELS = ['Qwen3-Embedding-0.6B', 'Qwen3-Embedding-4B', 'Qwen3-Embedding-8B', 'bge-m3',
          'embeddinggemma-300m', 'mmBERT-base', 'mmBERT-base_center', 'mmBERT-small',
          'mmBERT-small_center', 'xlm-roberta-base', 'xlm-roberta-base_center',
          'xlm-roberta-large', 'xlm-roberta-large_center']
TYPES_SHEET = ['流畅综合', '单独定义', '简单综合', '单独例句（均值池化）']
TYPES_FILE = ['流畅综合', '单独定义', '简单综合', '单独例句_均值池化']
NOTE = ('口径：无norm + cooc off-diagonal min-max 到 [0,1] + λ=0.1（blend = 0.1×相似度 + '
        '0.9×cooc_minmax）；指标为 Top-1 树的 F1（与原表 F1_top1 同口径）；random 行 = '
        '5 个种子的均值；cooccur_only 为 λ=0 纯共现、与 minmax 无关，直接沿用原表；'
        'case5/case8 缺 单独例句_均值池化 相似度矩阵，对应 13 格留空；case6/12 不做。')


def preflight():
    if XLSX.with_name(f'~${XLSX.name}').exists():
        sys.exit(f'文件被占用（存在锁文件 ~${XLSX.name}）：请先关闭 WPS/Excel')
    holding = subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip()
    if holding:
        sys.exit(f'文件被占用：\n{holding}\n请先关闭 WPS/Excel')
    print('预检通过：无进程占用、无锁文件')


def snapshot(wb, skip=()):
    out = {}
    for name in wb.sheetnames:
        if name in skip:
            continue
        ws = wb[name]
        out[name] = ([tuple(ws.cell(r, c).value for c in range(1, ws.max_column + 1))
                      for r in range(1, ws.max_row + 1)],
                     {str(m) for m in ws.merged_cells.ranges})
    return out


def main():
    preflight()
    df = pd.read_csv(SRC)
    ok = df[df.status == 'ok']
    assert len(ok) == 554, df.status.value_counts()
    fused = ok[ok.source == 'embsim_L融合'].set_index(['case', 'model', 'emb_type'])['F1_top1']
    random = ok[ok.source == 'random'].groupby('case')['F1_top1'].mean()
    char = ok[ok.source == 'char_sim'].set_index('case')['F1_top1']

    if not XLSX.stat().st_mode & 0o200:
        XLSX.chmod(XLSX.stat().st_mode | 0o600)
        print('已解除只读')
    stamp = datetime.now().strftime('%Y%m%dT%H%M%S')
    backup = XLSX.with_name(f'{XLSX.stem}.bak-{stamp}{XLSX.suffix}')
    backup.write_bytes(XLSX.read_bytes())
    print('备份 →', backup)

    wb = load_workbook(XLSX)
    before = snapshot(wb)
    old = wb[TEMPLATE]
    cooccur = {}
    last = None
    for r in range(2, old.max_row + 1):
        if old.cell(r, 1).value is not None:
            last = int(float(old.cell(r, 1).value))
        if old.cell(r, 2).value == 'cooccur_only':
            cooccur[f'case{last}'] = old.cell(r, 5).value

    if SHEET in wb.sheetnames:
        wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET)
    ws.append(['case', '融合方式', '模型名称', '综合方式', 'F1'])
    for c in range(1, 6):
        ws.cell(1, c).font = Font(name='宋体', size=11)
    for letter, width in zip('ABCDE', [8, 17.6, 24, 23.6, 11.5]):
        ws.column_dimensions[letter].width = width

    want_rows, blanks = {}, 0
    for i, case in enumerate(CASES):
        start = ws.max_row + 1
        ws.append([int(case.replace('case', '')), 'cooccur_only', None, None, cooccur[case]])
        ws.append([None, 'random', None, None, round(float(random[case]), 10)])
        ws.append([None, 'char_sim', None, None, round(float(char[case]), 10)])
        for j, model in enumerate(MODELS):
            for k, (type_file, type_sheet) in enumerate(zip(TYPES_FILE, TYPES_SHEET)):
                key = (case, model, type_file)
                value = fused.get(key, np.nan)
                ws.append([None,
                           'embsim_L融合' if (j == 0 and k == 0) else None,
                           model if k == 0 else None,
                           type_sheet,
                           None if pd.isna(value) else round(float(value), 10)])
                if pd.isna(value):
                    blanks += 1
        assert ws.max_row - start + 1 == 55, (case, ws.max_row - start + 1)
        ws.merge_cells(start_row=start, start_column=1, end_row=ws.max_row, end_column=1)
        ws.cell(start, 1).alignment = Alignment(vertical='center')
    assert ws.max_row == 551 and blanks == 26, (ws.max_row, blanks)
    ws.cell(1, 7).value = NOTE
    ws.cell(1, 7).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['G'].width = 70
    wb.save(XLSX)
    print(f'已写入「{SHEET}」{ws.max_row} 行 × {ws.max_column} 列（{blanks} 格留空）→ {XLSX.name}')

    # ---------- 校验 ----------
    new_wb = load_workbook(XLSX)
    errs = []
    after = snapshot(new_wb, skip=(SHEET,))
    for name, payload in before.items():
        if name == SHEET:
            continue
        if name not in after:
            errs.append(f'sheet 丢失：{name}')
        elif after[name] != payload:
            errs.append(f'sheet「{name}」内容或合并区发生变化')
    nws = new_wb[SHEET]
    if (nws.max_row, nws.max_column) != (551, 7):
        errs.append(f'新 sheet 尺寸 {(nws.max_row, nws.max_column)} != (551, 7)')
    expect_merges = {f'A{2 + 55 * i}:A{56 + 55 * i}' for i in range(10)}
    if {str(m) for m in nws.merged_cells.ranges} != expect_merges:
        errs.append('新 sheet case 列合并区不符')
    row = 2
    for case in CASES:
        if nws.cell(row, 1).value != int(case.replace('case', '')):
            errs.append(f'{case} case 值错位')
        if abs(float(nws.cell(row, 5).value) - float(round(cooccur[case], 10))) > 0:
            errs.append(f'{case} cooccur 值不符')
        if abs(float(nws.cell(row + 1, 5).value) - round(float(random[case]), 10)) > 0:
            errs.append(f'{case} random 值不符')
        if abs(float(nws.cell(row + 2, 5).value) - round(float(char[case]), 10)) > 0:
            errs.append(f'{case} char_sim 值不符')
        for j, model in enumerate(MODELS):
            for k, type_file in enumerate(TYPES_FILE):
                r = row + 3 + j * 4 + k
                if nws.cell(r, 3).value not in (None, model) or nws.cell(r, 4).value != TYPES_SHEET[k]:
                    errs.append(f'第 {r} 行标签不符')
                value = fused.get((case, model, type_file), np.nan)
                got = nws.cell(r, 5).value
                if pd.isna(value):
                    if got is not None:
                        errs.append(f'第 {r} 行应留空，实际 {got}')
                elif got is None or abs(float(got) - round(float(value), 10)) > 0:
                    errs.append(f'第 {r} 行 {got} != {round(float(value), 10)}')
        row += 55
    print('\n校验：' + ('全部通过 ✓' if not errs else f'{len(errs)} 项失败'))
    for e in errs[:10]:
        print('  ✗', e)
    if errs:
        print('还原备份…')
        XLSX.write_bytes(backup.read_bytes())
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
