"""把 case8 的三条 blend 行更新为与 app.py /api/process-excel 一致的结果。

依据：ARR_Oct/main_evaluation/rerun_case8/evaluation_summary.csv（用当前 main.py+SMM.py 重跑 case8），
已验证与 app 端点 Top-1 全指标一致。

写入范围（仅「主实验」sheet）：
  - case8 的 共现_blend_Jaccard_0.1 / 共现_blend_Dice_0.2 / 共现_blend_Emb_0.4 三行 × 5 块 × 10 指标
  - 平均（除去case6和12）块同名三行（按 10 个 case 的表内值重算）
其余任何单元格必须零变化，否则脚本报错退出。
"""
from __future__ import annotations

import contextlib
import io
import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(__file__).resolve().parents[4] / '实验记录_完整.xlsx'
FRESH_CSV = ROOT / 'ARR_Oct/main_evaluation/rerun_case8/evaluation_summary.csv'
CASE = 'case8'
CASE_ID = 8
INCLUDE_IDS = [1, 2, 3, 4, 5, 7, 8, 9, 10, 11]

METRICS = ['coverage', 'productivity', 'deg_mean', 'deg_std', 'weight_sum',
           'acc', 'prec', 'recall', 'F1', 'total_time_s']
BLOCKS = [(1, 'value', 3), (5, 'mean', 13), (5, 'max', 23),
          (10, 'mean', 33), (10, 'max', 43)]
METHOD2LABEL = {'Jaccard': '共现_blend_Jaccard_0.1',
                'Dice': '共现_blend_Dice_0.2',
                'Embeds': '共现_blend_Emb_0.4'}
LABEL2METHOD = {v: k for k, v in METHOD2LABEL.items()}


def cell_value(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    v = float(v)
    if abs(v - round(v)) < 1e-12:
        return int(round(v))
    return round(v, 10)


def preflight():
    if subprocess.run(['lsof', str(XLSX)], capture_output=True, text=True).stdout.strip():
        sys.exit(f'目标文件正被占用（WPS/Excel 打开中）：{XLSX}')
    locks = list(XLSX.parent.glob('~$*实验记录*'))
    if locks:
        print(f'[警告] 发现锁文件 {locks}，继续前请确认文件未被占用')


def find_block_start(ws, wanted):
    starts = [r for r in range(1, ws.max_row + 1)
              if ws.cell(r, 2).value == '纯共现（NAACL）']
    for r in starts:
        if str(ws.cell(r, 1).value) == str(wanted):
            return r
    sys.exit(f'未找到 case={wanted} 块')


def label_rows(ws, start):
    rows = {}
    for r in range(start, start + 11):
        rows[ws.cell(r, 2).value] = r
    if len(rows) != 11:
        sys.exit(f'块行数异常：{start} 得到 {len(rows)} 行')
    return rows


def read_row(ws, r):
    out = {}
    for k, st, col0 in BLOCKS:
        for i, m in enumerate(METRICS):
            out[(k, st, m)] = ws.cell(r, col0 + i).value
    return out


def app_check():
    """断言：app 端点 Top-1 / Top-5(mean,max) 与重跑结果一致（时间列除外）。"""
    sys.path.insert(0, str(ROOT))
    from app import app as flask_app

    data = pd.read_excel(ROOT / 'data_pure' / f'{CASE}.xlsx', sheet_name='Sheet1').astype(object)
    label = pd.read_excel(ROOT / 'data_pure' / f'{CASE}.xlsx', sheet_name='Sheet2').astype(object)
    base = dict(data=data.where(pd.notna(data), None).to_dict('records'),
                label=label.where(pd.notna(label), None).to_dict('records'))
    csv_text = (ROOT / 'ARR_Oct/LLM_similarity/case8/Qwen3-Embedding-8B'
                / 'function_similarity_matrix_单独定义.csv').read_text(encoding='utf-8')
    fresh = pd.read_csv(FRESH_CSV)
    errors = []
    for method, calc, w in [('Jaccard', 'J', 0.1), ('Dice', 'D', 0.2), ('Embeds', 'L', 0.4)]:
        payload = dict(base, calc_type=calc, blend_weight=w,
                       llm_sim_csv=csv_text if method == 'Embeds' else None)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            resp = flask_app.test_client().post(
                '/api/process-excel', data=json.dumps(payload, ensure_ascii=False),
                content_type='application/json')
        if resp.status_code != 200:
            sys.exit(f'app 端点 {method} 返回 {resp.status_code}: {resp.get_json()}')
        em = [g['evaluation_metric'] for g in resp.get_json()['graph_data']
              if not g.get('is_gt', False)]
        em = em[:5]
        if len(em) < 5:
            sys.exit(f'app 只返回 {len(em)} 个候选：{method}')
        for k, st in [(1, 'value'), (5, 'mean'), (5, 'max')]:
            row = fresh[(fresh.method == method) & (fresh.top_k == k)
                        & (fresh.statistic == st)]
            if row.empty:
                sys.exit(f'重跑结果缺少 {method} {k} {st}')
            row = row.iloc[0]
            for m in METRICS:
                if m == 'total_time_s':
                    continue
                if k == 1:
                    app_v = em[0][m]
                else:
                    vals = [e[m] for e in em]
                    app_v = float(np.mean(vals)) if st == 'mean' else float(np.max(vals))
                if abs(float(row[m]) - app_v) > 1e-9:
                    errors.append((method, k, st, m, float(row[m]), app_v))
    if errors:
        print('app 与重跑结果不一致：')
        for e in errors:
            print('   ', e)
        sys.exit('中止：app 与重跑不一致，未写入')
    print(f'app 一致性校验通过：3 方法 × Top1/Top5mean/Top5max × 9 指标全部一致')


def main():
    preflight()
    app_check()

    fresh = pd.read_csv(FRESH_CSV)
    wb = load_workbook(XLSX)
    ws = wb['主实验']
    header = [ws.cell(2, c).value for c in range(3, 53)]
    expected = [('time' if m == 'total_time_s' else m) for m in METRICS] * 5
    if header != expected:
        sys.exit(f'表头与预期不符：{header}')

    # 平均块行号
    avg_start = None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is not None and '平均' in str(a):
            avg_start = r
    if avg_start is None:
        sys.exit('未找到平均块')

    case_start = find_block_start(ws, CASE_ID)
    case_rows = label_rows(ws, case_start)
    avg_rows = label_rows(ws, avg_start)
    case_blocks = [(i, find_block_start(ws, i)) for i in INCLUDE_IDS]

    # 用未改动的「纯共现（NAACL）」行验证均值公式（10 case 算术平均）
    sanity = '纯共现（NAACL）'
    sanity_rows = label_rows(ws, avg_start)
    for k, st, col0 in BLOCKS:
        for i, m in enumerate(METRICS):
            c = col0 + i
            vals = [float(ws.cell(label_rows(ws, start)[sanity], c).value)
                    for _, start in case_blocks]
            got = float(ws.cell(sanity_rows[sanity], c).value)
            if abs(float(np.mean(vals)) - got) > 1e-9:
                sys.exit(f'均值公式自检失败 {k}/{st}/{m}: 算出 {np.mean(vals)} 表内 {got}')
    print('均值公式自检通过（纯共现（NAACL）行 10 case 均值与表内一致）')

    backup = XLSX.with_name(f"{XLSX.stem}.bak-{datetime.now():%Y%m%dT%H%M%S}.xlsx")
    shutil.copy2(XLSX, backup)
    print(f'备份：{backup}')

    before = {(r, c): ws.cell(r, c).value
              for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)}

    changed = []
    for method, label in METHOD2LABEL.items():
        r = case_rows[label]
        for k, st, col0 in BLOCKS:
            row = fresh[(fresh.method == method) & (fresh.top_k == k)
                        & (fresh.statistic == st)].iloc[0]
            for i, m in enumerate(METRICS):
                new = cell_value(row[m])
                c = col0 + i
                old = ws.cell(r, c).value
                if old != new:
                    changed.append((f'{label}', f'{k}/{st}/{m}', old, new))
                ws.cell(r, c).value = new

        # 平均块：按表内 10 个 case 的当前值重算
        ra = avg_rows[label]
        for k, st, col0 in BLOCKS:
            for i, m in enumerate(METRICS):
                c = col0 + i
                vals = []
                for cid, start in case_blocks:
                    v = ws.cell(label_rows(ws, start)[label], c).value
                    if v is None:
                        sys.exit(f'case{cid} {label} {k}/{st}/{m} 为空，无法求均值')
                    vals.append(float(v))
                new = cell_value(float(np.mean(vals)))
                old = ws.cell(ra, c).value
                if old != new:
                    changed.append(('平均（除去case6和12）', f'{label} {k}/{st}/{m}', old, new))
                ws.cell(ra, c).value = new

    wb.save(XLSX)
    print(f'写入 {len(changed)} 格变化')

    # ---- 全表零变化校验 ----
    wb2 = load_workbook(XLSX, data_only=False)
    ws2 = wb2['主实验']
    allowed = set()
    for label in METHOD2LABEL.values():
        allowed.update((case_rows[label], c) for c in range(3, 53))
        allowed.update((avg_rows[label], c) for c in range(3, 53))
    bad = []
    for (r, c), v0 in before.items():
        v1 = ws2.cell(r, c).value
        if (r, c) in allowed:
            continue
        if v0 != v1:
            bad.append((r, c, v0, v1))
    if bad:
        print(f'非目标格被改动 {len(bad)} 处：{bad[:10]}')
        shutil.copy2(backup, XLSX)
        sys.exit('已回滚')

    # 目标格必须等于重跑结果 / 与平均重算一致
    check = []
    for method, label in METHOD2LABEL.items():
        r = case_rows[label]
        for k, st, col0 in BLOCKS:
            row = fresh[(fresh.method == method) & (fresh.top_k == k)
                        & (fresh.statistic == st)].iloc[0]
            for i, m in enumerate(METRICS):
                exp = cell_value(row[m])
                got = ws2.cell(r, col0 + i).value
                if got != exp:
                    check.append((label, k, st, m, got, exp))
        ra = avg_rows[label]
        for k, st, col0 in BLOCKS:
            for i, m in enumerate(METRICS):
                c = col0 + i
                vals = [float(ws2.cell(label_rows(ws2, start)[label], c).value)
                        for _, start in case_blocks]
                exp = cell_value(float(np.mean(vals)))
                got = ws2.cell(ra, c).value
                if got != exp:
                    check.append(('avg', label, k, st, m, got, exp))
    if check:
        print(f'校验失败 {len(check)}：{check[:10]}')
        shutil.copy2(backup, XLSX)
        sys.exit('已回滚')

    print(f'校验通过：非目标 {len(before) - len(allowed)} 格零变化；目标格与重跑/均值完全一致')
    print('\n变化明细：')
    for c in changed:
        print('   ', c)


if __name__ == '__main__':
    main()
